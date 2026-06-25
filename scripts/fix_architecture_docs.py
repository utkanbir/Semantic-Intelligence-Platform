"""Targeted fixes for SIP architecture docx updates."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

ARCH = Path(__file__).resolve().parents[1] / "architecture"


def insert_paragraph_after(anchor: Paragraph, text: str) -> Paragraph:
    new_p = deepcopy(anchor._element)
    anchor._element.addnext(new_p)
    p = Paragraph(new_p, anchor._parent)
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)
    return p


def set_text(p: Paragraph, text: str) -> None:
    if p.runs:
        for r in p.runs:
            r.text = ""
        p.runs[0].text = text
    else:
        p.add_run(text)


def replace_status_block(doc: Document, heading: str, statuses: list[str]) -> None:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() != heading:
            continue
        j = i + 1
        while j < len(doc.paragraphs) and doc.paragraphs[j].text.strip() != "Suggested statuses:":
            j += 1
        if j >= len(doc.paragraphs):
            return
        set_text(doc.paragraphs[j], "Suggested statuses (SIP Asset Catalog v1 is authoritative):")
        k = j + 1
        for status in statuses:
            while k < len(doc.paragraphs) and doc.paragraphs[k].text.strip() == "":
                k += 1
            if k < len(doc.paragraphs):
                set_text(doc.paragraphs[k], status)
                k += 1
        return


def fix_domain_model() -> None:
    path = ARCH / "SIP Domain Model v1.docx"
    doc = Document(path)

    if not any(p.text.strip() == "├─ ontology_namespace" for p in doc.paragraphs):
        for p in doc.paragraphs:
            if p.text.strip() == "├─ metadata_domain":
                anchor = p
                for field in [
                    "├─ ontology_namespace",
                    "├─ agent_namespace",
                    "├─ product_registry_namespace",
                    "├─ agent_registry_namespace",
                ]:
                    anchor = insert_paragraph_after(anchor, field)
                break

    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "Provisioned":
            set_text(p, "Versioned")
        if "ApplicationWorkspace represents logical isolation, not dedicated infrastructure" in p.text:
            set_text(
                p,
                "ApplicationWorkspace represents the logical isolation boundary of an Application Workspace, including infrastructure and semantic/runtime registry namespaces. It stores isolation metadata only; it does not manage namespace contents.",
            )

    if not any("Which ontology namespace should be used?" in p.text for p in doc.paragraphs):
        for p in doc.paragraphs:
            if p.text.strip() == "Which metadata catalog domain should be used?":
                anchor = p
                for q in [
                    "Which ontology namespace should be used?",
                    "Which agent namespace should be used?",
                    "Which product registry namespace should be used?",
                    "Which agent registry namespace should be used?",
                ]:
                    anchor = insert_paragraph_after(anchor, q)
                break

    replace_status_block(doc, "Blueprint Status", ["Draft", "Review", "Approved", "Versioned", "Retired"])
    replace_status_block(
        doc,
        "PublishedDataProduct Status",
        ["Draft", "Certified", "Published", "Versioned", "Retired"],
    )
    replace_status_block(
        doc,
        "AgentDefinition Status",
        ["Draft", "Approved", "Active", "Versioned", "Retired"],
    )

    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Asset Status":
            for j in range(i + 1, min(i + 8, len(doc.paragraphs))):
                if doc.paragraphs[j].text.strip() == "Suggested statuses:":
                    set_text(
                        doc.paragraphs[j],
                        "Suggested statuses (per asset_type; SIP Asset Catalog v1 is authoritative):",
                    )
                    insert_paragraph_after(
                        doc.paragraphs[j],
                        "Align status values with SIP Asset Catalog v1 lifecycle states for each asset type.",
                    )
                    break
            break

    if not any("ARR-002" in p.text for p in doc.paragraphs):
        doc.add_paragraph(
            "Lifecycle authority (ARR-002): SIP Asset Catalog v1 is authoritative for asset lifecycle states. Domain Model status enums implement Asset Catalog semantics."
        )

    doc.save(path)
    print("Fixed:", path.name)


def fix_architecture_spec() -> None:
    path = ARCH / "SIP Architecture Specification.docx"
    doc = Document(path)

    for p in doc.paragraphs:
        if p.text.startswith("Provisioning creates namespace registrations"):
            set_text(
                p,
                "Provisioning creates namespace registrations and workspace metadata only. Applications remain blank workspaces (D-011, ARR-004). No domain ontologies, knowledge graphs, Published Data Products, or Agents are created automatically.",
            )

    texts = [p.text.strip() for p in doc.paragraphs]
    if "- Agent Namespace" not in texts and "- Ontology Namespace" in texts:
        for p in doc.paragraphs:
            if p.text.strip() == "- Ontology Namespace":
                insert_paragraph_after(p, "- Agent Namespace")
                break

    doc.save(path)
    print("Fixed:", path.name)


def fix_discovery_workflow() -> None:
    path = ARCH / "SIP Application Discovery Workflow v1.docx"
    doc = Document(path)

    if any("ARR-004" in p.text for p in doc.paragraphs):
        doc.save(path)
        print("Already fixed:", path.name)
        return

    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Phase 9 — Application Provisioning":
            for j in range(i, min(i + 20, len(doc.paragraphs))):
                if doc.paragraphs[j].text.strip() == "Application Workspace":
                    anchor = doc.paragraphs[j]
                    anchor = insert_paragraph_after(anchor, "Namespace registrations")
                    anchor = insert_paragraph_after(anchor, "Workspace metadata records")
                    insert_paragraph_after(
                        anchor,
                        "Note (ARR-004): Domain semantic assets are not created at this phase. Applications remain blank workspaces (D-011).",
                    )
                    break
            break

    doc.save(path)
    print("Fixed:", path.name)


def fix_user_journeys() -> None:
    path = ARCH / "SIP Core User Journeys v1.docx"
    doc = Document(path)

    for p in doc.paragraphs:
        if "11. SIP provisions" in p.text and "ARR-004" not in p.text:
            set_text(
                p,
                '11. SIP provisions:    - Application Workspace    - Namespace registrations    - Workspace metadata records    ("Initial Assets" = system-level records only; no domain ontologies, products, or agents — ARR-004, D-011)',
            )

    doc.save(path)
    print("Fixed:", path.name)


if __name__ == "__main__":
    fix_domain_model()
    fix_architecture_spec()
    fix_discovery_workflow()
    fix_user_journeys()
